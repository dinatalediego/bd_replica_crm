"""Import the governed, aggregate workbook baseline without treating it as ML evidence."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path


# Official codes confirmed against analytics.absorcion_inicio_proyecto.
# Fail closed for new projects until their identity has been checked in CORE.
PROJECT_CODES = {
    "Alicanto": "GY",
    "Capadocia": "CP",
    "Edificio Urbanzen": "EEUU",
    "Fenix": "FX",
    "Matera": "MT",
    "Modena": "MD",
    "Sialia": "SL",
    "Tizón y Bueno": "TZ",
    "Torre Nápoles": "NP",
}

REQUIRED = {
    ("1. Resumen comercial", "Meta Total Departamentos"): "meta",
    ("1. Resumen comercial", "Valor Total Colocado"): "colocado",
    ("1. Resumen comercial", "Gap a Meta"): "gap_reportado",
    ("4. Composición de ventas", "Valor Unidades Vendidas"): "vendido",
    ("4. Composición de ventas", "Valor Unidades Separadas"): "separado",
    ("5. Stock Por vender", "Valor Stock Disponible"): "stock_remanente",
    ("5. Stock Por vender", "Valor Unidades Disponible"): "stock_disponible",
}
BLOCKED = ("5. Stock Por vender", "Valor Unidades Bloqueadas")


@dataclass(frozen=True)
class Metric:
    code: str
    project: str
    block: str
    indicator: str
    metric_type: str
    original: str | None
    value: Decimal
    quantity: int | None


def _number(raw: object, label: str) -> Decimal:
    if raw is None or isinstance(raw, bool):
        raise ValueError(f"{label}: valor numérico ausente")
    try:
        result = Decimal(str(raw))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{label}: valor numérico inválido: {raw!r}") from exc
    if not result.is_finite():
        raise ValueError(f"{label}: valor numérico no finito")
    return result


def parse_rows(rows: list[tuple]) -> tuple[list[Metric], list[dict]]:
    metrics: list[Metric] = []
    by_project: dict[str, dict[tuple[str, str], Decimal]] = defaultdict(dict)
    for row in rows:
        block, indicator, project, original, metric_type, value, quantity = tuple(row[:7])
        if not (block and indicator and project) or block == "Bloque":
            continue
        project = str(project).strip()
        if project not in PROJECT_CODES:
            raise ValueError(f"Proyecto sin código CORE validado: {project!r}")
        key = (str(block).strip(), str(indicator).strip())
        if key in by_project[project]:
            raise ValueError(f"Indicador duplicado: {project} / {key}")
        number = _number(value, f"{project} / {key}")
        count = None
        if quantity is not None:
            integer = _number(quantity, f"{project} / cantidad {key}")
            if integer != integer.to_integral_value() or integer < 0:
                raise ValueError(f"{project} / cantidad {key}: debe ser entero no negativo")
            count = int(integer)
        by_project[project][key] = number
        metrics.append(Metric(PROJECT_CODES[project], project, *key,
                              str(metric_type or ""), str(original) if original is not None else None,
                              number, count))
    if not by_project:
        raise ValueError("La hoja Base no contiene proyectos")
    diagnostics: list[dict] = []
    for project, values in sorted(by_project.items()):
        missing = set(REQUIRED) - values.keys()
        if missing:
            raise ValueError(f"{project}: faltan indicadores: {sorted(missing)}")
        v = {name: values[key] for key, name in REQUIRED.items()}
        blocked = values.get(BLOCKED, Decimal(0))
        if v["meta"] <= 0 or any(v[k] < 0 for k in v if k != "meta") or blocked < 0:
            raise ValueError(f"{project}: meta debe ser positiva y los importes no negativos")
        stock_difference = v["stock_remanente"] - v["stock_disponible"] - blocked
        if abs(stock_difference) > 10:
            raise ValueError(f"{project}: stock no concilia (diferencia S/ {stock_difference})")
        composition_difference = v["colocado"] - v["vendido"] - v["separado"]
        if abs(composition_difference) > 10:
            raise ValueError(f"{project}: vendido + separado no concilia con colocado")
        gap_math = max(v["meta"] - v["colocado"], Decimal(0))
        diagnostics.append({
            "project": project, "code": PROJECT_CODES[project],
            "gap_reported": v["gap_reportado"], "gap_math": gap_math,
            "gap_difference": v["gap_reportado"] - gap_math,
            "stock_difference": stock_difference,
        })
    return metrics, diagnostics


def parse_workbook(path: Path) -> tuple[list[Metric], list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Base" not in wb.sheetnames:
            raise ValueError("Se requiere la hoja Base con columnas A:G")
        return parse_rows(list(wb["Base"].iter_rows(values_only=True)))
    finally:
        wb.close()
